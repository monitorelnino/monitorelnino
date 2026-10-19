import io, sys
from coletores_base import buscar
import openpyxl
b = buscar("https://ftp.ibge.gov.br/Perfil_Municipios/2020/Base_de_Dados/Base_MUNIC_2020.xlsx", timeout=180)
print("BYTES", len(b))
wb = openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True)
print("ABAS", wb.sheetnames)
for n in wb.sheetnames:
    if "ic" in n.lower() and "ion" in n.lower():
        ws = wb[n]
        print("DICIONARIO", n)
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            if i < 12: print(i, [str(c)[:70] for c in r if c is not None][:8])
        k = 0
        for r in ws.iter_rows(values_only=True):
            t = [str(c) for c in r if c is not None]
            if any(x.startswith("Mgrd") for x in t):
                k += 1
                if k < 14 or 180 < k < 190: print("MGRD", k, [x[:110] for x in t][:8])
        print("MGRD_LINHAS", k)
ws = wb["Gestão de riscos"]
rows = ws.iter_rows(values_only=True)
h = list(next(rows)); print("CABEC", len(h), h[:8], h[-5:])
print("MGRD_COLS", sum(1 for c in h if str(c).startswith("Mgrd")))
for i in range(3): print("LINHA", [str(c)[:12] for c in next(rows)][:12])
vals = {}
for r in rows:
    for c, v in zip(h, r):
        if str(c).startswith("Mgrd"): vals.setdefault(str(v), 0); vals[str(v)] += 1
print("VALORES_DISTINTOS", dict(sorted(vals.items(), key=lambda x: -x[1])[:25]))
