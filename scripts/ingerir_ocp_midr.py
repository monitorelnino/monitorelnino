#!/usr/bin/env python3
"""
scripts/ingerir_ocp_midr.py — a resposta do MIDR sobre a Operação Carro-Pipa
=============================================================================
Item 2 do `HANDOVER_preparacao_programatica_seca_05-10-2026.md`.

A FONTE
-------
Planilha "Dados OCP — Janeiro a Agosto — 2026" (SEI 7044288), entregue pelo MIDR em resposta a
pedido de acesso à informação, com o despacho SEI 7043337 de 29/09/2026. Os dois arquivos vivem no
repositório PRIVADO da editoria e **nunca** entram no repositório público: o que sai daqui é o
agregado por município, que é dado público, com a proveniência dita.

O hash é conferido antes da leitura. Planilha que não casa com o hash do `LEIAME.md` não é lida —
não por desconfiança do arquivo, mas porque o número publicado tem de poder ser refeito a partir de
um arquivo identificado.

AS TRÊS COISAS QUE ESTA INGESTÃO NÃO FAZ
-----------------------------------------
1. **Não soma médias por linha.** `Qtd. meses`, `Média de carradas`, `Média de pop` e
   `Média de vol água` são atributos do MUNICÍPIO, repetidos em cada linha dele. Somá-los por linha
   multiplicaria o município pelo número de portarias que ele tem — e 97 das 742 linhas são
   repetição de município.
2. **Não corrige anomalia.** `Qtd. meses = 9` em 10 linhas, num período de 8 meses, fica como está,
   registrado e contável. MG ausente embora tenha Semiárido fica registrado. A relação de
   "reconhecidos sem atendimento" prometida no despacho não veio, e isso é lacuna declarada.
3. **Não chama de "novos" os municípios com portaria a partir de 29/06/2026.** A planilha não
   separa reconhecimento de renovação, e o fato que ela sustenta é só este: "portaria com data a
   partir de 29/06/2026".

A CHAVE É O IBGE, NUNCA O NOME: os nomes vêm em maiúsculas, sem acento, com espaço sobrando.

USO
  python3 scripts/ingerir_ocp_midr.py --autoteste
  python3 scripts/ingerir_ocp_midr.py --dry-run
  python3 scripts/ingerir_ocp_midr.py
"""
import collections
import datetime as dt
import hashlib
import json
import pathlib
import sys

RAIZ = pathlib.Path(__file__).resolve().parent.parent
PLANILHA = pathlib.Path("C:/Users/User/Documents/MARE/robo-registro/notas/lai/respostas/MIDR/"
                        "ocp_dados_jan_ago_2026.xlsx")
HASH_ESPERADO = "6e4361b9b5750c59"
SAIDA = RAIZ / "data" / "programas_federais" / "ocp_2026.json"

PERIODO = "2026-01..2026-08"
MESES_DO_PERIODO = 8
FONTE = ("resposta do MIDR a pedido de acesso à informação, recebida em 05/10/2026 "
         "(planilha SEI 7044288; despacho SEI 7043337)")
RECEBIDO_EM = "2026-10-05"
INICIO_DO_CICLO = "2026-06-29"

COLUNAS = ("Ano", "UF", "Município", "IBGE", "Qtd. meses", "Média de carradas", "Média de pop",
           "Média de vol água", "Data da portaria", "Nº da portaria", "Situação de emergência",
           "Nº do D.O.U.", "Data do D.O.U.", "Processo")

# A planilha diz "Estiagem" ou "Seca". São as duas palavras da fonte, e ficam as duas: fundi-las
# apagaria uma distinção que o órgão faz.
SITUACOES = ("Estiagem", "Seca")


def texto(valor) -> str:
    """Valor de célula como texto limpo. Função pura."""
    return "" if valor is None else str(valor).strip()


def numero(valor):
    """Valor de célula como número, ou None. Função pura; não inventa zero no vazio."""
    s = texto(valor).replace(",", ".")
    if not s:
        return None
    try:
        f = float(s)
    except ValueError:
        return None
    return int(f) if f == int(f) else f


def data_iso(valor) -> str:
    """"dd/mm/aaaa" (ou data do Excel) → "aaaa-mm-dd". Função pura; "" no que não é data."""
    if isinstance(valor, dt.datetime):
        return valor.date().isoformat()
    if isinstance(valor, dt.date):
        return valor.isoformat()
    s = texto(valor)
    if len(s) == 10 and s[4] == "-" and s[7] == "-":
        return s
    partes = s.split("/")
    if len(partes) == 3 and len(partes[2]) == 4 and all(p.isdigit() for p in partes):
        return f"{partes[2]}-{partes[1].zfill(2)}-{partes[0].zfill(2)}"
    return ""


def chave_da_portaria(linha: dict) -> tuple:
    """A identidade de uma portaria: IBGE + número + data. Função pura.

    É esta tripla que o handover manda usar para deduplicar, e não o município: o mesmo município
    aparece em várias linhas, uma por portaria.
    """
    return (texto(linha.get("IBGE")), texto(linha.get("Nº da portaria")),
            data_iso(linha.get("Data da portaria")))


def agregar(linhas: list) -> dict:
    """{ibge: registro do município}, com as portarias distintas. Função pura.

    Os atributos do município (meses, médias) são lidos da PRIMEIRA linha dele; quando outra linha
    discorda, a divergência entra em `divergencias` do próprio município em vez de virar média de
    médias — média de médias não é o número de ninguém.
    """
    por_ibge = {}
    vistas = set()
    for linha in linhas or []:
        ibge = texto(linha.get("IBGE"))
        if not ibge:
            continue
        reg = por_ibge.get(ibge)
        if reg is None:
            reg = por_ibge[ibge] = {
                "ibge": ibge, "uf": texto(linha.get("UF")),
                "nome_na_fonte": " ".join(texto(linha.get("Município")).split()),
                "meses_atendidos": numero(linha.get("Qtd. meses")),
                "media_de_carradas": numero(linha.get("Média de carradas")),
                "media_de_populacao": numero(linha.get("Média de pop")),
                "media_de_volume_de_agua": numero(linha.get("Média de vol água")),
                "portarias": [], "divergencias": [],
            }
        for campo, coluna in (("meses_atendidos", "Qtd. meses"),
                              ("media_de_carradas", "Média de carradas"),
                              ("media_de_populacao", "Média de pop"),
                              ("media_de_volume_de_agua", "Média de vol água")):
            valor = numero(linha.get(coluna))
            if valor is not None and reg[campo] is not None and valor != reg[campo]:
                aviso = f"{coluna}: {reg[campo]} numa linha e {valor} em outra"
                if aviso not in reg["divergencias"]:
                    reg["divergencias"].append(aviso)

        chave = chave_da_portaria(linha)
        if chave in vistas:
            continue
        vistas.add(chave)
        reg["portarias"].append({
            "numero": texto(linha.get("Nº da portaria")),
            "data": data_iso(linha.get("Data da portaria")),
            "situacao_na_fonte": texto(linha.get("Situação de emergência")),
            "dou_numero": texto(linha.get("Nº do D.O.U.")),
            "dou_data": data_iso(linha.get("Data do D.O.U.")),
            "processo": texto(linha.get("Processo")),
        })
    for reg in por_ibge.values():
        reg["portarias"].sort(key=lambda p: (p["data"], p["numero"]))
    return por_ibge


def anomalias(linhas: list, por_ibge: dict) -> list:
    """O que a planilha tem de estranho, descrito e não consertado. Função pura."""
    fora = []
    nove = sorted({(texto(x.get("UF")), " ".join(texto(x.get("Município")).split()),
                    texto(x.get("IBGE")))
                   for x in linhas or []
                   if (numero(x.get("Qtd. meses")) or 0) > MESES_DO_PERIODO})
    if nove:
        linhas_nove = sum(1 for x in linhas
                          if (numero(x.get("Qtd. meses")) or 0) > MESES_DO_PERIODO)
        fora.append({
            "anomalia": "meses_atendidos_acima_do_periodo",
            "descricao": (f"`Qtd. meses` acima de {MESES_DO_PERIODO} em {linhas_nove} linha(s), "
                          f"{len(nove)} município(s), num período de {MESES_DO_PERIODO} meses"),
            "municipios": [{"uf": u, "nome_na_fonte": n, "ibge": i} for u, n, i in nove],
            "tratamento": "registrada, não corrigida",
        })
    ufs = sorted({r["uf"] for r in por_ibge.values() if r["uf"]})
    if "MG" not in ufs:
        fora.append({
            "anomalia": "uf_do_semiarido_ausente",
            "descricao": "Minas Gerais não aparece na planilha, embora tenha território no "
                         "Semiárido",
            "municipios": [], "tratamento": "registrada, não corrigida",
        })
    fora.append({
        "anomalia": "relacao_prometida_nao_entregue",
        "descricao": "o despacho SEI 7043337 anuncia a relação de municípios reconhecidos sem "
                     "atendimento, e ela não veio com a resposta",
        "municipios": [], "tratamento": "lacuna declarada",
    })
    return fora


def resumir(por_ibge: dict) -> dict:
    """Os números agregados, contados e não estimados. Função pura."""
    portarias = [(ibge, p) for ibge, r in por_ibge.items() for p in r["portarias"]]
    por_uf = collections.Counter(r["uf"] for r in por_ibge.values() if r["uf"])
    por_meses = collections.Counter(r["meses_atendidos"] for r in por_ibge.values()
                                    if r["meses_atendidos"] is not None)
    por_situacao = collections.Counter(p["situacao_na_fonte"] for _, p in portarias)
    return {
        "municipios": len(por_ibge),
        "ufs": sorted(por_uf),
        "municipios_por_uf": dict(sorted(por_uf.items())),
        "portarias_distintas": len(portarias),
        "municipios_por_meses_atendidos": {str(k): v for k, v in sorted(por_meses.items())},
        "portarias_por_situacao_na_fonte": dict(sorted(por_situacao.items())),
        # O handover é explícito: a planilha NÃO separa reconhecimento de renovação, então o fato
        # permitido é a data da portaria, e nada sobre ser o município "novo" no programa.
        "portarias_com_data_a_partir_do_inicio_do_ciclo": sum(
            1 for _, p in portarias if p["data"] and p["data"] >= INICIO_DO_CICLO),
        "inicio_do_ciclo": INICIO_DO_CICLO,
    }


def montar(linhas: list) -> dict:
    """O arquivo inteiro, a partir das linhas da planilha. Função pura — não lê e não escreve."""
    por_ibge = agregar(linhas)
    return {
        "programa": "Operação Carro-Pipa",
        "orgao": "Ministério da Integração e do Desenvolvimento Regional",
        "periodo": PERIODO,
        "fonte": FONTE,
        "recebido_em": RECEBIDO_EM,
        "linhas_na_fonte": len([x for x in linhas or [] if texto(x.get("IBGE"))]),
        "resumo": resumir(por_ibge),
        "anomalias": anomalias(linhas, por_ibge),
        "municipios": [por_ibge[k] for k in sorted(por_ibge)],
    }


def ler_planilha(caminho: pathlib.Path) -> list:
    """As linhas da planilha, conferido o hash. Lê disco; não escreve."""
    import openpyxl
    achado = hashlib.sha256(caminho.read_bytes()).hexdigest()
    if not achado.startswith(HASH_ESPERADO):
        raise SystemExit(f"{caminho.name}: sha256 começa em {achado[:16]} e o LEIAME diz "
                         f"{HASH_ESPERADO} — planilha não conferida, não foi lida")
    livro = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    folha = livro[livro.sheetnames[0]]
    cru = list(folha.iter_rows(values_only=True))
    cabecalho = [texto(c) for c in cru[0]]
    faltam = [c for c in COLUNAS if c not in cabecalho]
    if faltam:
        raise SystemExit(f"{caminho.name}: colunas ausentes: {', '.join(faltam)}")
    return [dict(zip(cabecalho, linha)) for linha in cru[1:]
            if any(c is not None for c in linha)]


def _autoteste() -> int:
    falhas, contados = [], []

    def ok(nome, cond):
        contados.append(nome)
        print(("  ✓ " if cond else "  ✗ ") + nome)
        if not cond:
            falhas.append(nome)

    ok("data brasileira vira ISO", data_iso("26/03/2026") == "2026-03-26")
    ok("data já em ISO passa", data_iso("2026-03-26") == "2026-03-26")
    ok("data de Excel vira ISO", data_iso(dt.datetime(2026, 3, 26)) == "2026-03-26")
    ok("o que não é data devolve vazio", data_iso("sem data") == "")
    ok("decimal com vírgula é número", numero("597,1") == 597.1)
    ok("inteiro não vira float", isinstance(numero("57"), int))
    ok("vazio não vira zero", numero("") is None)
    ok("nome com espaço sobrando é normalizado",
       agregar([{"IBGE": "1", "Município": "  AGUA  BRANCA "}])["1"]["nome_na_fonte"]
       == "AGUA BRANCA")

    # Duas linhas do MESMO município com portarias diferentes: um município, duas portarias, e os
    # atributos NÃO se somam — foi o erro que o handover nomeou.
    base = {"IBGE": "2700102", "UF": "AL", "Município": "AGUA BRANCA", "Qtd. meses": 6,
            "Média de carradas": 57, "Média de pop": 1288, "Média de vol água": 597.1,
            "Situação de emergência": "Estiagem", "Nº do D.O.U.": "59",
            "Data do D.O.U.": "27/03/2026", "Processo": "59051.046866/2026-51"}
    duas = [dict(base, **{"Nº da portaria": "1022", "Data da portaria": "26/03/2026"}),
            dict(base, **{"Nº da portaria": "1500", "Data da portaria": "30/06/2026"})]
    ag = agregar(duas)
    ok("duas portarias do mesmo município dão um município", len(ag) == 1)
    ok("as duas portarias ficam", len(ag["2700102"]["portarias"]) == 2)
    ok("os atributos do município não somam", ag["2700102"]["media_de_carradas"] == 57)
    ok("as portarias ficam em ordem de data",
       [p["numero"] for p in ag["2700102"]["portarias"]] == ["1022", "1500"])

    repetida = duas + [dict(duas[0])]
    ok("linha repetida não cria portaria",
       len(agregar(repetida)["2700102"]["portarias"]) == 2)

    divergente = [duas[0], dict(duas[1], **{"Qtd. meses": 7})]
    ok("atributo que discorda entre linhas vira divergência registrada",
       any("Qtd. meses" in d for d in agregar(divergente)["2700102"]["divergencias"]))
    ok("divergência não muda o valor publicado",
       agregar(divergente)["2700102"]["meses_atendidos"] == 6)

    r = resumir(agregar(duas))
    ok("a contagem de portarias é de portarias distintas", r["portarias_distintas"] == 2)
    ok("portaria a partir de 29/06 é contada pela data",
       r["portarias_com_data_a_partir_do_inicio_do_ciclo"] == 1)
    ok("o resumo não usa a palavra «novos»", "novo" not in json.dumps(r, ensure_ascii=False))

    nove = [dict(base, **{"Nº da portaria": "1", "Data da portaria": "01/02/2026",
                          "Qtd. meses": 9})]
    an = anomalias(nove, agregar(nove))
    ok("mês acima do período vira anomalia",
       any(a["anomalia"] == "meses_atendidos_acima_do_periodo" for a in an))
    ok("anomalia é registrada, não corrigida",
       agregar(nove)["2700102"]["meses_atendidos"] == 9)
    ok("MG ausente vira anomalia", any(a["anomalia"] == "uf_do_semiarido_ausente" for a in an))
    ok("MG presente não vira anomalia",
       not any(a["anomalia"] == "uf_do_semiarido_ausente"
               for a in anomalias(nove, {"1": {"uf": "MG", "portarias": []}})))
    ok("a relação prometida e não entregue é lacuna declarada",
       any(a["tratamento"] == "lacuna declarada" for a in an))

    inteiro = montar(duas)
    ok("o arquivo diz a fonte", "acesso à informação" in inteiro["fonte"])
    ok("o arquivo diz o período", inteiro["periodo"] == PERIODO)
    ok("linha sem IBGE não entra na contagem da fonte",
       montar(duas + [{"UF": "AL"}])["linhas_na_fonte"] == 2)

    import dis
    nomes = set()
    for nome_obj, obj in list(globals().items()):
        if nome_obj in ("_autoteste", "main", "ler_planilha"):
            continue
        codigo = getattr(obj, "__code__", None)
        if codigo is not None:
            nomes |= {i.argval for i in dis.get_instructions(codigo) if isinstance(i.argval, str)}
    ok("trava estrutural: as funções puras não escrevem",
       not ({"write_text", "write_bytes", "open"} & nomes))
    ok("trava estrutural: as funções puras não vão à rede",
       # `get` e `post` ficam FORA desta lista: `dict.get` aparece em toda função pura daqui, e
       # uma trava que acusa `dict.get` de ir à rede ensina a ignorá-la.
       not ({"urlopen", "requests", "httpx", "urllib3", "Session"} & nomes))

    print(("✗ AUTOTESTE: " + str(len(falhas)) + " falha(s)") if falhas
          else f"✓ AUTOTESTE OK — {len(contados)} casos, sem rede e sem escrita.")
    return 1 if falhas else 0


def main(argv: list) -> int:
    if "--autoteste" in argv:
        return _autoteste()
    if not PLANILHA.exists():
        print(f"✗ {PLANILHA} não existe nesta árvore — a planilha vive no repositório privado")
        return 1
    dados = montar(ler_planilha(PLANILHA))
    r = dados["resumo"]
    print(f"OCP {dados['periodo']}: {r['municipios']} municípios em {len(r['ufs'])} estados · "
          f"{r['portarias_distintas']} portarias distintas em {dados['linhas_na_fonte']} linhas · "
          f"{r['portarias_com_data_a_partir_do_inicio_do_ciclo']} com data a partir de "
          f"{INICIO_DO_CICLO}")
    for a in dados["anomalias"]:
        print(f"  ⚠ {a['anomalia']}: {a['descricao']} [{a['tratamento']}]")
    if "--dry-run" in argv:
        print("  (--dry-run: nada escrito)")
        return 0
    SAIDA.parent.mkdir(parents=True, exist_ok=True)
    SAIDA.write_text(json.dumps(dados, ensure_ascii=False, indent=1, sort_keys=False) + "\n",
                     encoding="utf-8", newline="\n")
    print(f"  → {SAIDA.relative_to(RAIZ)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
